"""批量导入个人投递记录，并与当前岗位库自动匹配。"""
from __future__ import annotations

import csv
import datetime as dt
import io
import itertools
import re
import uuid
from typing import Any

import openpyxl
from openpyxl.utils.datetime import from_excel
from sqlalchemy import desc
from sqlalchemy.orm import Session

from ..models import Application, ApplicationStage, Group, Job, User
from .stage_service import DEFAULT_APPLICATION_STAGES


HEADER_ALIASES = {
    "公司": "company",
    "公司名称": "company",
    "公司名字": "company",
    "企业": "company",
    "company": "company",
    "岗位": "title",
    "岗位名称": "title",
    "投递岗位": "title",
    "职位": "title",
    "职位名称": "title",
    "title": "title",
    "投递时间": "applied_at",
    "投递日期": "applied_at",
    "日期": "applied_at",
    "appliedat": "applied_at",
    "applied_at": "applied_at",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalized_header(value: Any) -> str:
    return re.sub(r"[\s_]+", "", _text(value).lower())


def normalized_application_key(company: Any, title: Any) -> tuple[str, str]:
    def normalize(value: Any) -> str:
        return re.sub(r"\s+", "", _text(value)).casefold()

    return normalize(company), normalize(title)


def parse_applied_at(value: Any) -> dt.datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, (int, float)) and 1 <= value <= 100000:
        try:
            parsed = from_excel(value)
            if isinstance(parsed, dt.datetime):
                return parsed.replace(tzinfo=None)
            if isinstance(parsed, dt.date):
                return dt.datetime.combine(parsed, dt.time.min)
        except (TypeError, ValueError, OverflowError):
            return None

    text = _text(value)
    if not text:
        return None
    normalized = (
        text.replace("年", "-")
        .replace("月", "-")
        .replace("日", "")
        .replace("/", "-")
        .replace(".", "-")
        .replace("T", " ")
        .replace("Z", "")
    )
    try:
        return dt.datetime.fromisoformat(normalized).replace(tzinfo=None)
    except ValueError:
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d", "%y-%m-%d"):
            try:
                return dt.datetime.strptime(normalized, fmt)
            except ValueError:
                continue
    return None


def _rows_from_matrix(matrix: list[list[Any]]) -> list[dict]:
    if not matrix:
        return []
    first = matrix[0]
    mapped = [HEADER_ALIASES.get(_normalized_header(value)) for value in first]
    has_header = "company" in mapped and "title" in mapped
    data_rows = matrix[1:] if has_header else matrix
    default_keys = ("company", "title", "applied_at")
    rows: list[dict] = []
    for values in data_rows:
        if not any(_text(value) for value in values):
            continue
        row: dict[str, Any] = {}
        for index, value in enumerate(values):
            key = mapped[index] if has_header and index < len(mapped) else (
                default_keys[index] if index < len(default_keys) else None
            )
            if key and key not in row:
                row[key] = value
        rows.append(row)
    return rows


def parse_application_import_file(content: bytes, filename: str) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        matrix = [
            list(row)
            for row in itertools.islice(worksheet.iter_rows(values_only=True), 1002)
        ]
        workbook.close()
        return _rows_from_matrix(matrix)
    if name.endswith(".csv"):
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV")
        sample = decoded[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        matrix = [
            list(row)
            for row in itertools.islice(csv.reader(io.StringIO(decoded), dialect), 1002)
        ]
        return _rows_from_matrix(matrix)
    raise ValueError("仅支持 .xlsx 或 .csv 文件")


def serializable_import_rows(rows: list[dict], limit: int = 1000) -> dict:
    output = []
    errors = []
    for index, row in enumerate(rows[:limit], start=1):
        company = _text(row.get("company"))
        title = _text(row.get("title"))
        applied_at = parse_applied_at(row.get("applied_at"))
        error = None
        if not company or not title:
            error = "公司名称和岗位名称不能为空"
        elif not applied_at:
            error = "投递时间不能为空或格式不正确"
        if error:
            errors.append({"row": index, "message": error})
        output.append({
            "_row": index,
            "company": company,
            "title": title,
            "applied_at": applied_at.isoformat(timespec="minutes") if applied_at else _text(row.get("applied_at")),
            "_error": error,
        })
    if len(rows) > limit:
        errors.append({"row": limit + 1, "message": f"单次最多导入 {limit} 行"})
    return {"rows": output, "errors": errors}


def import_application_rows(
    db: Session,
    user: User,
    group: Group,
    raw_rows: list[dict],
    limit: int = 1000,
) -> dict:
    prepared = serializable_import_rows(raw_rows, limit=limit)
    errors = list(prepared["errors"])

    existing_application_keys = {
        (*normalized_application_key(company, title), applied_at.date().isoformat())
        for company, title, applied_at in db.query(
            Application.company,
            Application.title,
            Application.applied_at,
        ).filter(Application.user_id == user.id).all()
        if applied_at
    }
    jobs = db.query(Job).filter(
        Job.group_id == group.id,
        Job.is_active.is_not(False),
    ).order_by(desc(Job.updated_at), desc(Job.id)).all()
    jobs_by_key: dict[tuple[str, str], Job] = {}
    for job in jobs:
        jobs_by_key.setdefault(normalized_application_key(job.company, job.title), job)

    accepted: list[dict] = []
    results: list[dict] = []
    duplicate_count = 0
    new_jobs: list[Job] = []
    new_job_keys: set[tuple[str, str]] = set()
    matched_jobs = 0

    for index, row in enumerate(prepared["rows"], start=1):
        source_row = int(row.get("_row") or index)
        if row.get("_error"):
            continue
        applied_at = parse_applied_at(row["applied_at"])
        company = row["company"]
        title = row["title"]
        normalized = normalized_application_key(company, title)
        duplicate_key = (*normalized, applied_at.date().isoformat())
        if duplicate_key in existing_application_keys:
            duplicate_count += 1
            results.append({
                "row": source_row,
                "company": company,
                "title": title,
                "applied_at": applied_at.date().isoformat(),
                "status": "duplicate",
                "message": "重复记录，已跳过",
            })
            continue
        existing_application_keys.add(duplicate_key)
        job = jobs_by_key.get(normalized)
        if job:
            if normalized in new_job_keys:
                job_status = "created"
            else:
                matched_jobs += 1
                job_status = "matched"
        else:
            job = Job(
                source="manual",
                source_id=f"g{group.id}-application-import-{uuid.uuid4().hex}",
                group_id=group.id,
                created_by_id=user.id,
                company=company,
                title=title,
                is_active=True,
            )
            db.add(job)
            new_jobs.append(job)
            new_job_keys.add(normalized)
            jobs_by_key[normalized] = job
            job_status = "created"
        accepted.append({
            "company": job.company,
            "title": job.title,
            "applied_at": applied_at,
            "job": job,
            "job_status": job_status,
            "row": source_row,
        })

    if new_jobs:
        db.flush()

    applications: list[Application] = []
    for item in accepted:
        app = Application(
            user_id=user.id,
            job_id=item["job"].id,
            company=item["company"],
            title=item["title"],
            status="已投递",
            current_stage="投递",
            applied_at=item["applied_at"],
        )
        db.add(app)
        applications.append(app)
        item["application"] = app
    if applications:
        db.flush()

    stages: list[ApplicationStage] = []
    for item in accepted:
        app = item["application"]
        for position, stage_name in enumerate(DEFAULT_APPLICATION_STAGES):
            stages.append(ApplicationStage(
                application_id=app.id,
                stage=stage_name,
                position=position,
                status="completed" if position == 0 else "pending",
                schedule_type="exact",
                completed_at=item["applied_at"] if position == 0 else None,
            ))
        results.append({
            "row": item["row"],
            "company": item["company"],
            "title": item["title"],
            "applied_at": item["applied_at"].date().isoformat(),
            "status": "imported",
            "job_status": item["job_status"],
            "message": "已匹配岗位库" if item["job_status"] == "matched" else "已新增岗位并记录投递",
        })
    if stages:
        db.add_all(stages)
    db.commit()

    return {
        "total": len(raw_rows),
        "imported": len(accepted),
        "matched_jobs": matched_jobs,
        "created_jobs": len(new_jobs),
        "duplicates": duplicate_count,
        "failed": len(errors),
        "errors": errors,
        "items": sorted(results, key=lambda item: item["row"]),
    }
